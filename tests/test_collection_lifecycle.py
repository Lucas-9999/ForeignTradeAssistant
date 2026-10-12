"""Operational safety tests for paused collection sources and Fumeng intake."""

from __future__ import annotations

import sqlite3
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest


def test_source_registry_hides_paused_sources():
    from collector.cli import build_source_registry

    settings = SimpleNamespace(
        COLLECTION_INDUSTRY_ENABLED=True,
        COLLECTION_XUNKEXING_ENABLED=False,
        FB_COLLECTION_ENABLED=False,
        COLLECTION_LINKEDIN_ENABLED=False,
    )

    registry = build_source_registry(settings)

    assert "industry" in registry
    assert "xunkexing" not in registry
    assert "facebook" not in registry
    assert "linkedin" not in registry


def test_paused_scheduled_collection_returns_without_creating_task(monkeypatch):
    import config
    from collector.scheduler import _run_scheduled_collection

    monkeypatch.setattr(config, "COLLECTION_XUNKEXING_ENABLED", False)
    _run_scheduled_collection()


def test_paused_scheduler_status_has_no_collection_next_run(monkeypatch):
    import config
    import collector.scheduler as scheduler

    monkeypatch.setattr(config, "COLLECTION_XUNKEXING_ENABLED", False)
    scheduler.stop_scheduler(wait=False)
    try:
        scheduler.start_scheduler()
        status = scheduler.get_scheduler_status()
        assert status["next_run"]["paused"] is True
        assert status["next_run"]["scheduled_at"] is None
        assert "daily_collection" not in {job["id"] for job in status["jobs"]}
    finally:
        scheduler.stop_scheduler(wait=False)


def test_followup_schedule_can_be_disabled_without_disabling_imap(monkeypatch):
    """P13 can keep bounce monitoring on while holding auto-follow-ups."""
    import config
    import collector.scheduler as scheduler

    monkeypatch.setattr(config, "COLLECTION_XUNKEXING_ENABLED", False)
    monkeypatch.setattr(config, "EMAIL_SCHEDULE_ENABLED", False)
    monkeypatch.setattr(config, "ZOHO_BATCH_ENABLED", False)
    monkeypatch.setattr(config, "FOLLOWUP_SCHEDULE_ENABLED", False)
    monkeypatch.setattr(config, "IMAP_CHECK_ENABLED", True)
    scheduler.stop_scheduler(wait=False)
    try:
        scheduler.start_scheduler()
        job_ids = {job["id"] for job in scheduler.get_scheduler_status()["jobs"]}
        assert "daily_inbox_check" in job_ids
        assert "daily_followup_check" not in job_ids
    finally:
        scheduler.stop_scheduler(wait=False)


def test_email_schedule_start_date_blocks_catchup_until_the_requested_day(monkeypatch):
    import config
    import collector.scheduler as scheduler

    monkeypatch.setattr(config, "EMAIL_SCHEDULE_START_DATE", "2026-08-25")

    assert scheduler._email_schedule_is_active(datetime(2026, 8, 24, 10, 0)) is False
    assert scheduler._email_schedule_is_active(datetime(2026, 8, 25, 10, 0)) is True


def test_default_batch_excludes_archived_xunkexing(tmp_path, monkeypatch):
    import config
    from crm.db import build_database
    from emailer.batch import send_batch

    db_path = tmp_path / "archive.db"
    build_database(db_path).close()
    monkeypatch.setattr(config, "DB_PATH", db_path)
    monkeypatch.setattr(config, "OUTREACH_ARCHIVED_SOURCES", frozenset({"xunkexing"}))
    monkeypatch.setattr("utils.email_verify.has_mx_record", lambda _domain: True)
    conn = sqlite3.connect(str(db_path))
    conn.execute(
        "INSERT INTO leads (company_name, email, source, status, discovered_at) "
        "VALUES ('Archived XK', 'archived@example.org', 'xunkexing', 'new', datetime('now'))"
    )
    conn.execute(
        "INSERT INTO leads (company_name, email, source, status, discovered_at) "
        "VALUES ('Manual', 'manual@example.org', 'manual', 'new', datetime('now'))"
    )
    conn.commit()
    conn.close()

    result = send_batch(limit=5, dry_run=True)

    assert [detail["company"] for detail in result["details"]] == ["Manual"]


def test_fumeng_preview_reads_alias_headers_without_writing(tmp_path):
    from collector.imports.fumeng import preview_workbook

    path = tmp_path / "fumeng.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["公司名称", "联系人", "职位", "电子邮箱", "国家/地区", "电话"])
    sheet.append(["Acme Fasteners", "Alice", "Buyer", "alice@acme.example", "US", "+1 555 0100"])
    sheet.append(["Duplicate", "Alice 2", "Buyer", "ALICE@ACME.EXAMPLE", "US", None])
    sheet.append(["No Email", "Bob", None, None, "DE", None])
    workbook.save(path)
    workbook.close()

    preview = preview_workbook(path)

    assert preview.total_rows == 3
    assert preview.valid_rows == 1
    assert preview.missing_email_rows == 1
    assert preview.duplicate_email_rows == 1
    assert preview.records[0]["company_name"] == "Acme Fasteners"
    assert preview.records[0]["email"] == "alice@acme.example"


def test_fumeng_preview_requires_an_email_column(tmp_path):
    from collector.imports.fumeng import FumengContractError, preview_workbook

    path = tmp_path / "invalid.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["联系人", "国家"])
    workbook.active.append(["Alice", "US"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(FumengContractError, match="email"):
        preview_workbook(path)


def test_fumeng_preview_cli_is_read_only(tmp_path, capsys):
    from tools.preview_fumeng_excel import main

    path = tmp_path / "fumeng.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Email"])
    workbook.active.append(["buyer@example.org"])
    workbook.save(path)
    workbook.close()

    assert main(["--file", str(path)]) == 0
    assert "no database writes" in capsys.readouterr().out


def test_industry_collector_extracts_public_plain_text_email():
    from bs4 import BeautifulSoup
    from collector.sources.industry import IndustryCollector

    collector = IndustryCollector(delay=0)
    soup = BeautifulSoup(
        "<h1>Acme</h1><p>Sales: sales@acme-fasteners.example</p>", "lxml"
    )

    result = collector._parse_fields({"fields": {"company_name": "h1"}}, soup, "https://example.test")

    assert result["email"] == "sales@acme-fasteners.example"
